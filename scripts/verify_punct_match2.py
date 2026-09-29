"""补充检查：v2 JSONL 的 total_errors 分布、排序结构、errors 子字段。"""
import json
from collections import Counter
from openpyxl import load_workbook

JL = r'C:\hyt-agent\训练集-微调版提示词\标点符号错误_训练集_v2.jsonl'

jl = []
with open(JL, 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        d = json.loads(line)
        user_c = asst_c = ''
        for m in d['messages']:
            if m['role'] == 'user':
                user_c = m['content']
            elif m['role'] == 'assistant':
                asst_c = m['content']
        payload = json.loads(user_c.split('[INPUT_PAYLOAD]', 1)[1].strip())
        jl.append({'ref': payload['reference'], 'dc': payload['detection_content'], 'a': json.loads(asst_c)})

# 1. total_errors 分布
pos = [j for j in jl if j['a']['has_error']]
print('有错样本 total_errors 分布:', dict(Counter(j['a']['total_errors'] for j in pos)))
print('有错样本 len(errors) 与 total_errors 一致:', sum(1 for j in pos if len(j['a']['errors']) == j['a']['total_errors']), '/', len(pos))
print()

# 2. errors 子字段
subfields = Counter()
for j in pos:
    for e in j['a']['errors']:
        for k in e:
            subfields[k] += 1
print('errors 子字段出现次数:', dict(subfields))
print()

# 3. JSONL 排序：是否也是前481负后580正
labels = [j['a']['has_error'] for j in jl]
first_pos = labels.index(True)
print(f'第一条正样本位置: {first_pos} (0-based)')
neg_only_prefix = all(not b for b in labels[:481])
pos_only_suffix = all(b for b in labels[481:])
print(f'前481条全为负: {neg_only_prefix} | 482之后全为正: {pos_only_suffix}')
print()

# 4. 按序对比：负样本行的 dc 与模板行关系
TPL = r'C:\hyt-agent\数据集\完整数据集标点符号.xlsx'
wb = load_workbook(TPL)
ws = wb['Sheet1']
header = [c.value for c in ws[1]]
tpl_rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

neg_jl = [j for j in jl if not j['a']['has_error']]
neg_tpl = [t for t in tpl_rows if not t['is_real_error']]
pos_jl = [j for j in jl if j['a']['has_error']]
pos_tpl = [t for t in tpl_rows if t['is_real_error']]

# 负样本：按序 dc 是否等于模板 error_content
neg_seq = sum(1 for j, t in zip(neg_jl, neg_tpl) if j['dc'] == t['error_content'])
print(f'负样本按序 dc==模板error_content: {neg_seq}/481')
# 负样本按序 reference 的题型是否等于模板 title
import re
def parse_ref(ref):
    s = re.search(r'【学科】(.+)', ref)
    t = re.search(r'【试卷标题】(.+)', ref)
    q = re.search(r'【题型或其它信息】(.+)', ref)
    return (s.group(1).strip() if s else None, t.group(1).strip() if t else None, q.group(1).strip() if q else None)

tt_match = sum(1 for j, t in zip(neg_jl, neg_tpl) if parse_ref(j['ref'])[2] == t['title'])
print(f'负样本按序 题型==模板title: {tt_match}/481')
print()

# 正样本：按序 dc 与模板 error_content 关系
pos_seq_eq = sum(1 for j, t in zip(pos_jl, pos_tpl) if j['dc'] == t['error_content'])
pos_seq_diff = len(pos_jl) - pos_seq_eq
print(f'正样本按序 dc==模板error_content: {pos_seq_eq} | 不同: {pos_seq_diff}')
# 正样本按序 题型匹配
ptt = sum(1 for j, t in zip(pos_jl, pos_tpl) if parse_ref(j['ref'])[2] == t['title'])
print(f'正样本按序 题型==模板title: {ptt}/580')
print()

# 5. 正样本中 dc 不同的行，看看是多错还是单错
diff_rows = [(j, t) for j, t in zip(pos_jl, pos_tpl) if j['dc'] != t['error_content']]
print(f'dc 不同的正样本: {len(diff_rows)} 行, 其中多错(total_errors>1): {sum(1 for j, t in diff_rows if j["a"]["total_errors"] > 1)}')
same_rows = [(j, t) for j, t in zip(pos_jl, pos_tpl) if j['dc'] == t['error_content']]
print(f'dc 相同的正样本: {len(same_rows)} 行, 其中多错: {sum(1 for j, t in same_rows if j["a"]["total_errors"] > 1)}')
print()

# 6. 负样本 dc 不同的行数
neg_diff = [(j, t) for j, t in zip(neg_jl, neg_tpl) if j['dc'] != t['error_content']]
print(f'dc 不同的负样本: {len(neg_diff)} 行')
for j, t in neg_diff[:3]:
    print(f'  JSONL dc: {j["dc"][:60]!r}')
    print(f'  模板 ec:  {str(t["error_content"])[:60]!r}')
    print()

# 7. 模板 title 取值全集 vs JSONL 题型全集
print('JSONL 题型全集:', sorted(set(parse_ref(j['ref'])[2] for j in jl)))
print('模板 title 全集:', sorted(set(str(t['title']) for t in tpl_rows)))
