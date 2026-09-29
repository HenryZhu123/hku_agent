"""核查：has_error=false 但 errors 非空的行，FP 是评估系统原判还是我自己统计的。"""
from openpyxl import load_workbook

MY = r'C:\hyt-agent\训练集-微调版提示词\train_v1.1\模型表现\train_v1.1_错题汇总_checkpoint-175.xlsx'
SRC = r'C:\Users\ASUS\Downloads\train_v1.1_智能评估.xlsx'

# 1. 在我生成的错题Excel里找截图对应的行（描述含"当前内容属于非选择题"）
wb = load_workbook(MY, read_only=True)
ws = wb['全部错题']
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('【我生成的错题Excel】列名:', header[:18], '...')
mine = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(header, r))
    if '当前内容属于非选择题' in str(d.get('gold标准错误(correction)', '')):
        mine.append(d)
print(f'我Excel里描述匹配"当前内容属于非选择题"的行数: {len(mine)}')
print('前三行的关键字段:')
for d in mine[:3]:
    print(f'  id={d.get("id")} | 出错分类={d.get("出错分类")} | has_error={d.get("模型has_error")} | is_real_error={d.get("is_real_error")} | gold={d.get("题目错误类型")}')
wb.close()
print()

# 2. 检查源评估文件：这些行 FP/TN/TP/FN 列是怎么标的
wb = load_workbook(SRC, read_only=True)
ws = wb['聚合检测']
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
src = {str(r[header.index('id')]): dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[header.index('id')]}
wb.close()

print('=== 源评估文件原判对比 ===')
for d in mine[:5]:
    sid = str(d.get('id'))
    if sid not in src:
        continue
    s = src[sid]
    fp = s.get('FP')
    tn = s.get('TN')
    tp = s.get('TP')
    fn = s.get('FN')
    he = s.get('has_error')
    real = s.get('is_real_error')
    gtype = s.get('error_type')
    erraw = str(s.get('errors'))[:60]
    # 评估系统原判
    sys_mark = '?' 
    if fp: sys_mark = 'FP'
    if tn: sys_mark = 'TN'
    if tp: sys_mark = 'TP'
    if fn: sys_mark = 'FN'
    my_mark = d.get('出错分类')
    print(f'  id={sid} | gold={gtype} | is_real={real} | has_error={he}')
    print(f'    评估系统原判: {sys_mark} (TP={tp} TN={tn} FP={fp} FN={fn})')
    print(f'    我自己的分类: {my_mark}')
    print(f'    errors: {erraw}')
    print()

# 3. 统计：源评估文件这些行里 FP 列有多少真值
src_match = []
for sid, s in src.items():
    if '当前内容属于非选择题' in str(s.get('errors','')):
        src_match.append(s)
print(f'=== 源评估文件内有同样描述的行数: {len(src_match)} ===')
print('它们在评估系统里的标记分布:')
from collections import Counter
print(' ', dict(Counter('FP' if s.get('FP') else ('TN' if s.get('TN') else ('TP' if s.get('TP') else ('FN' if s.get('FN') else '无'))) for s in src_match)))
print('它们的 has_error 列分布:', dict(Counter(str(s.get('has_error')) for s in src_match)))
