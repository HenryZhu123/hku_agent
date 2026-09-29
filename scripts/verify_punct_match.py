"""验证 JSONL v2 与模板是否逐行对应（按序 + 按 detection_content 匹配）。"""
import json
from collections import Counter
from openpyxl import load_workbook

TPL = r'C:\hyt-agent\数据集\完整数据集标点符号.xlsx'
JL = r'C:\hyt-agent\训练集-微调版提示词\标点符号错误_训练集_v2.jsonl'

wb = load_workbook(TPL)
ws = wb['Sheet1']
header = [c.value for c in ws[1]]
tpl_rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

jl = []
with open(JL, 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        d = json.loads(line)
        sys_c = user_c = asst_c = ''
        for m in d['messages']:
            if m['role'] == 'system':
                sys_c = m['content']
            elif m['role'] == 'user':
                user_c = m['content']
            else:
                asst_c = m['content']
        payload = json.loads(user_c.split('[INPUT_PAYLOAD]', 1)[1].strip())
        a = json.loads(asst_c)
        jl.append({'reference': payload['reference'], 'detection_content': payload['detection_content'], 'assistant': a})

print(f'JSONL: {len(jl)} 条 | 模板: {len(tpl_rows)} 行')
print()

# 1. 按顺序比对 detection_content vs error_content
seq_match = sum(1 for j, t in zip(jl, tpl_rows) if j['detection_content'] == t['error_content'])
print(f'按顺序 detection_content == error_content: {seq_match}/{len(jl)}')
seq_pos_match = sum(1 for j, t in zip(jl, tpl_rows) if j['detection_content'] == t['content'])
print(f'按顺序 detection_content == content: {seq_pos_match}/{len(jl)}')

# 正负标签按序一致性
lab_match = sum(1 for j, t in zip(jl, tpl_rows) if bool(j['assistant']['has_error']) == bool(t['is_real_error']))
print(f'按顺序 has_error == is_real_error: {lab_match}/{len(jl)}')
print()

# 2. 按内容匹配（多重集）
from collections import Counter as Cnt
jl_dc = Cnt(j['detection_content'] for j in jl)
tpl_ec = Cnt(t['error_content'] for t in tpl_rows)
tpl_ct = Cnt(t['content'] for t in tpl_rows)
common_ec = sum((jl_dc & tpl_ec).values())
common_ct = sum((jl_dc & tpl_ct).values())
print(f'detection_content 与模板 error_content 交集: {common_ec}')
print(f'detection_content 与模板 content 交集: {common_ct}')
print()

# 3. 匹配上的行中，正负标签是否一致
ec_set = set(tpl_ec)
agree = disagree = 0
for j in jl:
    if j['detection_content'] in ec_set:
        # 找模板中对应行
        for t in tpl_rows:
            if t['error_content'] == j['detection_content']:
                if bool(j['assistant']['has_error']) == bool(t['is_real_error']):
                    agree += 1
                else:
                    disagree += 1
                break
print(f'内容匹配行中正负标签: 一致{agree} 不一致{disagree}')
print()

# 4. 对匹配的正样本，比对 correction 关系：模板 content 是否等于 JSONL detection_content 按 errors 修正后的结果
def apply_fix(text, errs):
    """用 errors 的 original_text->correction 修正文本（每处只换第一次出现）。"""
    out = text
    for e in errs:
        ot = e.get('original_text', '')
        co = e.get('correction', '')
        if ot and ot != co:
            idx = out.find(ot)
            if idx >= 0:
                out = out[:idx] + co + out[idx + len(ot):]
    return out

checked = fixed_ok = 0
for j, t in zip(jl, tpl_rows):
    if j['detection_content'] == t['error_content'] and j['assistant']['has_error']:
        checked += 1
        if apply_fix(j['detection_content'], j['assistant']['errors']) == t['content']:
            fixed_ok += 1
print(f'正样本修正验证: apply_fix(detection_content, errors) == 模板content: {fixed_ok}/{checked}')

# 5. corrected_content 与 error_content / content 的关系（有错样本）
cc_eq_ec = sum(1 for t in tpl_rows if t['is_real_error'] and t['corrected_content'] is not None and t['corrected_content'] == t['error_content'])
cc_eq_ct = sum(1 for t in tpl_rows if t['is_real_error'] and t['corrected_content'] is not None and t['corrected_content'] == t['content'])
cc_none = sum(1 for t in tpl_rows if t['is_real_error'] and t['corrected_content'] is None)
print(f'有错样本 corrected_content: 等于error_content {cc_eq_ec} | 等于content {cc_eq_ct} | None {cc_none}')

# 6. correction 列格式统计
import re
fmt_ok = sum(1 for t in tpl_rows if t['is_real_error'] and re.match(r'^原文：.+ -> 错误：.+$', str(t['correction']), re.S))
print(f'有错样本 correction 符合"原文：X -> 错误：Y"格式: {fmt_ok}/580')

# 7. 模板 school/file_name 与 reference 的映射验证（用匹配行验证）
import re as _re
map_ok = map_fail = 0
for j, t in zip(jl, tpl_rows):
    if j['detection_content'] == t['error_content']:
        ref = j['reference']
        m_s = _re.search(r'【学科】(.+)', ref)
        m_t = _re.search(r'【试卷标题】(.+)', ref)
        m_q = _re.search(r'【题型或其它信息】(.+)', ref)
        sch = m_s.group(1).strip() if m_s else None
        fn = (m_t.group(1).strip().lstrip('# ').strip() + '.pdf') if m_t else None
        ti = m_q.group(1).strip() if m_q else None
        if sch == t['school'] and fn == t['file_name'] and ti == t['title']:
            map_ok += 1
        else:
            map_fail += 1
            if map_fail <= 3:
                print(f'  映射不一致: ref学科={sch!r} vs school={t["school"]!r} | fn={fn!r} vs {t["file_name"]!r} | ti={ti!r} vs {t["title"]!r}')
print(f'reference->school/file_name/title 映射验证: 一致{map_ok} 不一致{map_fail}')
