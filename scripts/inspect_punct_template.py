"""检查模板XLSX结构和JSONL格式。"""
import json
from openpyxl import load_workbook

TPL = r'C:\hyt-agent\数据集\完整数据集标点符号.xlsx'
JL = r'C:\hyt-agent\训练集-微调版提示词\标点符号错误_训练集_v2.jsonl'

# 1. 模板结构
wb = load_workbook(TPL)
print('=== 模板 XLSX ===')
print('工作表:', wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print(f'主表: {ws.max_row} 行 x {ws.max_column} 列')
print('表头:', [c.value for c in ws[1]])
print('冻结窗格:', ws.freeze_panes)
print('自动筛选:', ws.auto_filter.ref)
# 列宽
widths = {k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}
print('列宽:', widths)
# 表头样式
h = ws.cell(row=1, column=1)
print('表头字体:', h.font.bold, h.font.size, '| 填充:', h.fill.fgColor.rgb if h.fill and h.fill.fgColor else None)
print('表头对齐:', h.alignment.horizontal, h.alignment.vertical, 'wrap:', h.alignment.wrap_text)

# 抽样前3行数据，看每列类型和内容
print()
for r in range(2, 5):
    print(f'--- 行{r} ---')
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        cell = ws.cell(row=r, column=c)
        vs = str(v)[:80] if v is not None else 'None'
        print(f'  列{c}({ws.cell(row=1,column=c).value}) 类型={type(v).__name__} wrap={cell.alignment.wrap_text} : {vs}')
wb.close()
print()

# 2. JSONL 结构
print('=== JSONL ===')
with open(JL, 'r', encoding='utf-8') as f:
    lines = f.readlines()
print('总行数:', len(lines))
d = json.loads(lines[0])
print('顶层键:', list(d.keys()))
for m in d['messages']:
    print(f"--- {m['role']} 前300字 ---")
    print(m['content'][:300])
    print()

# 统计
n_pos = n_neg = 0
parse_fail = 0
for i, line in enumerate(lines):
    line = line.strip()
    if not line:
        continue
    d = json.loads(line)
    asst = None
    for m in d['messages']:
        if m['role'] == 'assistant':
            asst = m['content']
            break
    try:
        a = json.loads(asst)
        if a.get('has_error'):
            n_pos += 1
        else:
            n_neg += 1
    except Exception:
        parse_fail += 1
print(f'has_error=true: {n_pos} | has_error=false: {n_neg} | 解析失败: {parse_fail}')
