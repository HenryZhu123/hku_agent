# -*- coding: utf-8 -*-
"""生成 train_v1.1 智能评估审校错误汇总 Excel 报告"""
import json
from collections import Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = r'C:\Users\ASUS\Downloads\train_v1.1_智能评估.xlsx'
OUT = r'C:\hyt-agent\模型输出\train_v1.1_智能评估-审校错误汇总.xlsx'

wb = load_workbook(SRC, read_only=True)
ws = wb['聚合检测']
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

def parse_errors(v):
    if not v:
        return []
    try:
        el = json.loads(v) if isinstance(v, str) else v
        return el if isinstance(el, list) else []
    except Exception:
        return []

def errs_brief(v, maxlen=120):
    errs = parse_errors(v)
    if not errs:
        return ''
    parts = []
    for e in errs[:3]:
        ot = str(e.get('original_text', ''))[:15]
        co = str(e.get('correction', ''))[:15]
        parts.append(f'"{ot}"→"{co}"')
    s = '; '.join(parts)
    if len(errs) > 3:
        s += f' 等{len(errs)}条'
    return s[:maxlen]

def classify_fp(errs):
    """对 FP 行的报错内容分类"""
    cats = []
    for e in errs:
        ot = str(e.get('original_text', ''))
        co = str(e.get('correction', ''))
        desc = str(e.get('description', ''))
        if ot and ot == co:
            cats.append('幻觉错别字（原文=修改）')
        elif '非选择题' in desc or '无法提取选项' in desc or '不包含选项' in desc or '无法进行选项结构' in desc or '不包含需要检查的选项' in desc or '与要求的选项结构不匹配' in desc or '与检测任务要求' in desc or '与需要检测的选择题' in desc or '未提供可供检查的选项' in desc:
            cats.append('非选择题硬套选项结构检测')
        elif '题目类型识别错误' in desc or '作答形式' in str(co) or '改写' in str(co) or '题目类型' in str(co):
            cats.append('题型误判/要求改写题面')
        elif '错别字或拼写错误' in desc:
            cats.append('错别字误报（知识性改写或正确词被判错）')
        elif '选项标号序列' in desc or '未按A' in desc or '标号' in desc:
            cats.append('选项标号序列误报')
        elif '重复选项' in desc or '完全相同' in desc:
            cats.append('重复选项误报')
        elif '缺失' in desc or '不完整' in desc:
            cats.append('内容缺失类误报')
        else:
            cats.append('其他: ' + desc[:30])
    return ' | '.join(dict.fromkeys(cats)) if cats else ''

# 分类汇总数据
fp_rows = [r for r in rows if r.get('FP')]
fn_rows = [r for r in rows if r.get('FN')]
unmarked = [r for r in rows if r.get('TP') is None and r.get('TN') is None and r.get('FN') is None and r.get('FP') is None]
sk_pos = [r for r in unmarked if r.get('is_real_error') == 1]

out_wb = Workbook()

# ============ Sheet 1: 总览 ============
ws1 = out_wb.active
ws1.title = '总览'
header_font = Font(bold=True, size=11)
fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')

tp = sum(1 for r in rows if r.get('TP'))
tn = sum(1 for r in rows if r.get('TN'))
overview = [
    ['审校错误汇总总览', ''],
    ['', ''],
    ['评估文件', SRC],
    ['数据来源', 'train_v1.1 训练集智能评估（Qwen3.5-9B-finetune-v1.1-merged, checkpoint-175）'],
    ['总行数', len(rows)],
    ['（对照）训练集原始行数', 1500],
    ['缺失未评估行数', 16],
    ['', ''],
    ['混淆矩阵（行级）', ''],
    ['TP（正确检出）', tp],
    ['TN（正确判无错）', tn],
    ['FP（误报行数）', len(fp_rows)],
    ['FN（漏报行数）', len(fn_rows)],
    ['路由跳过行（判断题过滤）', len(unmarked)],
    ['其中：跳过的正样本（实际漏报）', len(sk_pos)],
    ['', ''],
    ['官方指标', ''],
    ['precision', 0.822], ['recall', 0.98], ['f1', 0.894], ['f2', 0.944], ['accuracy', 0.893], ['fpr', 0.182],
    ['', ''],
    ['审校出错合计', f'{len(fp_rows)}误报 + {len(fn_rows)}漏报 + {len(sk_pos)}跳过正样本 = {len(fp_rows)+len(fn_rows)+len(sk_pos)} 行'],
    ['', ''],
    ['误报按 gold 类型', ''],
    ['错别字', sum(1 for r in fp_rows if r.get('error_type') == '错别字')],
    ['选项结构错误', sum(1 for r in fp_rows if r.get('error_type') == '选项结构错误')],
    ['题目与题型不一致', sum(1 for r in fp_rows if r.get('error_type') == '题目与题型不一致')],
    ['', ''],
    ['漏报按 gold 类型', ''],
    ['错别字', sum(1 for r in fn_rows if r.get('error_type') == '错别字')],
    ['选项结构错误', sum(1 for r in fn_rows if r.get('error_type') == '选项结构错误')],
    ['题目与题型不一致', sum(1 for r in fn_rows if r.get('error_type') == '题目与题型不一致')],
    ['', ''],
    ['主要出错模式', ''],
    ['模式1：判断题路由过滤过激', f'{len(unmarked)}条判断题被整体跳过（理由：判断题待判陈述属于考点），其中{len(sk_pos)}条含真实错误从未审校'],
    ['模式2：幻觉错别字', '44处"错字"的原文与修改完全相同（如 动摇→动摇、解离度→解离度、DSC→DSC）'],
    ['模式3：知识性改写当错别字', '把正确的术语改写当错字报出（己糖激酶→葡萄糖激酶、甾类激素→甾体激素、1000V→1000kV、钢坯→钢版）'],
    ['模式4：非选择题硬套选项结构检测', '名词解释/简答/论述/作文题被报"无法提取选项标号"类无意义错误'],
    ['模式5：题型误判', '对正常题目报题型不一致并要求改写（判断题→论述题、简答题→问答题、名词解释→简答题）'],
    ['模式6：选项标号序列苛求', '对合法的D/E/F起始编号或格式差异报"未按A开始连续排列"'],
    ['', ''],
    ['缺失的16条评估行ID', ''],
    ['全部为负样本', 'TRAIN-V1.1-0190/0442/0457/0486/0560/0624/0662/0909/1007/1022/1119/1270/1388/1411/1471/1499（转换去重所致，含1119与1471内容重复）'],
]
for i, (a, b) in enumerate(overview, 1):
    ws1.cell(row=i, column=1, value=a).font = header_font
    ws1.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical='top')
ws1.column_dimensions['A'].width = 32
ws1.column_dimensions['B'].width = 95
ws1.freeze_panes = 'A2'

# ============ Sheet 2: FP误报明细 ============
ws2 = out_wb.create_sheet('FP误报明细')
cols2 = ['id', 'gold错误类型', 'is_real_error', 'FP处数', '误报模式', '模型报错内容', 'gold说明(gold_reason)', '题面摘要', '评估reason']
for c, name in enumerate(cols2, 1):
    cell = ws2.cell(row=1, column=c, value=name)
    cell.font = header_font
    cell.fill = fill
for i, r in enumerate(sorted(fp_rows, key=lambda x: (x.get('error_type') or '', str(x.get('id')))), 2):
    errs = parse_errors(r.get('errors'))
    ws2.cell(row=i, column=1, value=r.get('id'))
    ws2.cell(row=i, column=2, value=r.get('error_type'))
    ws2.cell(row=i, column=3, value=r.get('is_real_error'))
    ws2.cell(row=i, column=4, value=r.get('FP'))
    ws2.cell(row=i, column=5, value=classify_fp(errs))
    ws2.cell(row=i, column=6, value=errs_brief(r.get('errors'), 200))
    ws2.cell(row=i, column=7, value=str(r.get('gold_reason') or ''))
    ws2.cell(row=i, column=8, value=str(r.get('error_content') or '')[:100].replace('\n', ' / '))
    ws2.cell(row=i, column=9, value=str(r.get('reason') or '')[:250])
    for c in range(1, 10):
        ws2.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical='top')
for c, w in zip(range(1, 10), [18, 14, 10, 8, 34, 40, 28, 40, 45]):
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:{get_column_letter(len(cols2))}{len(fp_rows)+1}'

# ============ Sheet 3: FN漏报明细 ============
ws3 = out_wb.create_sheet('FN漏报明细')
cols3 = ['id', 'gold错误类型', 'gold错误数', 'gold错误内容', '模型输出(has_error)', '模型报错数', '题面摘要', '评估reason']
for c, name in enumerate(cols3, 1):
    cell = ws3.cell(row=1, column=c, value=name)
    cell.font = header_font
    cell.fill = fill
for i, r in enumerate(fn_rows, 2):
    gold_errs = parse_errors(r.get('correction'))
    model_errs = parse_errors(r.get('errors'))
    ws3.cell(row=i, column=1, value=r.get('id'))
    ws3.cell(row=i, column=2, value=r.get('error_type'))
    ws3.cell(row=i, column=3, value=r.get('error_cnt'))
    ws3.cell(row=i, column=4, value=errs_brief(r.get('correction'), 200))
    ws3.cell(row=i, column=5, value=str(r.get('has_error')))
    ws3.cell(row=i, column=6, value=len(model_errs))
    ws3.cell(row=i, column=7, value=str(r.get('error_content') or '')[:100].replace('\n', ' / '))
    ws3.cell(row=i, column=8, value=str(r.get('reason') or '')[:250])
    for c in range(1, 9):
        ws3.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical='top')
for c, w in zip(range(1, 9), [18, 14, 10, 40, 14, 10, 45, 45]):
    ws3.column_dimensions[get_column_letter(c)].width = w
ws3.freeze_panes = 'A2'

# ============ Sheet 4: 跳过正样本明细 ============
ws4 = out_wb.create_sheet('路由跳过正样本')
cols4 = ['id', 'gold错误类型', 'gold错误数', 'gold错误内容', '题面摘要', 'gold_reason', '跳过原因(failure_reason)']
for c, name in enumerate(cols4, 1):
    cell = ws4.cell(row=1, column=c, value=name)
    cell.font = header_font
    cell.fill = fill
for i, r in enumerate(sorted(sk_pos, key=lambda x: x.get('id') or ''), 2):
    ws4.cell(row=i, column=1, value=r.get('id'))
    ws4.cell(row=i, column=2, value=r.get('error_type'))
    ws4.cell(row=i, column=3, value=r.get('error_cnt'))
    ws4.cell(row=i, column=4, value=errs_brief(r.get('correction'), 200))
    ws4.cell(row=i, column=5, value=str(r.get('error_content') or '')[:100].replace('\n', ' / '))
    ws4.cell(row=i, column=6, value=str(r.get('gold_reason') or ''))
    ws4.cell(row=i, column=7, value=str(r.get('failure_reason') or ''))
    for c in range(1, 8):
        ws4.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical='top')
for c, w in zip(range(1, 8), [18, 14, 10, 40, 45, 30, 30]):
    ws4.column_dimensions[get_column_letter(c)].width = w
ws4.freeze_panes = 'A2'
ws4.auto_filter.ref = f'A1:{get_column_letter(len(cols4))}{len(sk_pos)+1}'

out_wb.save(OUT)
print(f'报告已生成: {OUT}')
print(f'FP误报: {len(fp_rows)} 行 / FN漏报: {len(fn_rows)} 行 / 跳过正样本: {len(sk_pos)} 行')

# 幻觉错别字统计
halluc_cnt = sum(1 for r in fp_rows for e in parse_errors(r.get('errors'))
                 if str(e.get('original_text', '')) and str(e.get('original_text', '')) == str(e.get('correction', '')))
print(f'其中幻觉错别字(原文=修改): {halluc_cnt} 处')
