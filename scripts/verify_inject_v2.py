"""用改进后的 inject（anchor 前后文定位）重新验证 580 正样本的自洽性。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from punct_repair import inject

from openpyxl import load_workbook

JL = r'C:\hyt-agent\训练集-微调版提示词\标点符号错误_训练集_v2.jsonl'
OUT = r'C:\hyt-agent\数据集\完整数据集标点符号_v2.xlsx'

jl = []
with open(JL, 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        d = json.loads(line)
        user_c = asst_c = ''
        for m in d['messages']:
            if m['role'] == 'user':
                user_c = m['content']
            elif m['role'] == 'assistant':
                asst_c = m['content']
        payload = json.loads(user_c.split('[INPUT_PAYLOAD]', 1)[1].strip())
        jl.append({'dc': payload['detection_content'], 'a': json.loads(asst_c)})

wb = load_workbook(OUT)
ws = wb['Sheet1']
header = [c.value for c in ws[1]]
rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

ok = loc_fail = res_fail = 0
fails = []
for r, j in zip(rows, jl):
    if not r['is_real_error']:
        continue
    out = r['content']
    failed = False
    for e in j['a']['errors']:
        out = inject(out, e)
        if out is None:
            loc_fail += 1
            fails.append((r['id'], '定位失败'))
            failed = True
            break
    if failed:
        continue
    if out == r['error_content']:
        ok += 1
    else:
        res_fail += 1
        fails.append((r['id'], '结果不等'))

print(f'正向注入验证: 成功 {ok} | 定位失败 {loc_fail} | 结果不等 {res_fail}')
for f in fails[:8]:
    print(' ', f)
