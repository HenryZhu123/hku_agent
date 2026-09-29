"""检查模板中有错样本的各列填法（非read_only模式）。"""
import json
from collections import Counter
from openpyxl import load_workbook

TPL = r'C:\hyt-agent\数据集\完整数据集标点符号.xlsx'

wb = load_workbook(TPL)
ws = wb['Sheet1']
header = [c.value for c in ws[1]]
rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

pos = [r for r in rows if r['is_real_error']]
neg = [r for r in rows if not r['is_real_error']]
print(f'模板: 总{len(rows)}行, 有错{len(pos)}, 无错{len(neg)}')
print()

for r in pos[:2]:
    print(f'===== 有错样本 id={r["id"]} =====')
    for k in header:
        v = str(r[k])[:300] if r[k] is not None else 'None'
        print(f'  {k}: {v}')
    print()

neg_nonnull = Counter(k for r in neg for k in header if r[k] is not None)
print('无错样本非空列分布:', dict(neg_nonnull))
pos_nonnull = Counter(k for r in pos for k in header if r[k] is not None)
print('有错样本非空列分布:', dict(pos_nonnull))
print()
print('有错 error_cnt 分布:', dict(Counter(r['error_cnt'] for r in pos)))
print('无错 error_cnt 分布:', dict(Counter(r['error_cnt'] for r in neg)))
print('error_type 分布:', dict(Counter(r['error_type'] for r in rows)))
print('error_detailed_type 取值(有错):', dict(Counter(r['error_detailed_type'] for r in pos)))
print('is_content_same 分布(有错):', dict(Counter(r['is_content_same'] for r in pos)))
print('is_content_same 分布(无错):', dict(Counter(r['is_content_same'] for r in neg)))
print('model_name 分布:', dict(Counter(r['model_name'] for r in rows)))
print()

# correction 字段格式：是 JSON 串还是文本？抽3个看
print('=== correction 格式抽样 ===')
for r in pos[:3]:
    c = r['correction']
    print(f'id={r["id"]} 类型={type(c).__name__}:')
    print(' ', str(c)[:400])
    # 尝试解析为JSON
    try:
        j = json.loads(c) if isinstance(c, str) else None
        print('  -> 可解析为JSON:', type(j).__name__ if j is not None else 'N/A')
    except Exception as e:
        print('  -> 不是JSON:', str(e)[:60])
    print()

# design_reason / corrected_content 抽样
print('=== design_reason 抽样 ===')
for r in pos[:2]:
    print(f'id={r["id"]}:', str(r['design_reason'])[:200])
print()
print('=== corrected_content 抽样 ===')
for r in pos[:2]:
    print(f'id={r["id"]}:', str(r['corrected_content'])[:200])
