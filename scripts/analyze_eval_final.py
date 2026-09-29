# -*- coding: utf-8 -*-
"""最终分析：标记共存矩阵、FP=2行、幻觉错别字模式、跳过正样本的gold错误、缺失行"""
import json
from collections import Counter
from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\ASUS\Downloads\train_v1.1_智能评估.xlsx', read_only=True)
ws = wb['聚合检测']
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()

def parse_errors(v):
    if not v:
        return []
    try:
        el = json.loads(v) if isinstance(v, str) else v
        return el if isinstance(el, list) else []
    except Exception:
        return []

# ---------- 1. 标记共存矩阵 ----------
print('=== 行级标记组合 ===')
combo = Counter()
for r in rows:
    flags = []
    if r.get('TP'): flags.append('TP')
    if r.get('TN'): flags.append('TN')
    if r.get('FN'): flags.append('FN')
    if r.get('FP'): flags.append(f"FP{r.get('FP')}" if r.get('FP') != 1 else 'FP')
    combo['+'.join(flags) if flags else '无标记'] += 1
for k, v in combo.most_common():
    print(f'  {k}: {v}')

# ---------- 2. TP且FN的行 ----------
tp_fn = [r for r in rows if r.get('TP') and r.get('FN')]
print(f'\nTP+FN 共存行: {len(tp_fn)}')
for r in tp_fn:
    print(f"  id={r.get('id')} TP={r.get('TP')} FN={r.get('FN')} FP={r.get('FP')}")

# ---------- 3. FP=2 的行 ----------
fp2 = [r for r in rows if r.get('FP') == 2]
print(f'\n=== FP=2 的行（{len(fp2)} 条） ===')
for r in fp2:
    errs = parse_errors(r.get('errors'))
    print(f"id={r.get('id')} [{r.get('error_type')}] is_real={r.get('is_real_error')} 模型报{len(errs)}条:")
    for e in errs:
        print(f"    '{str(e.get('original_text',''))[:25]}' → '{str(e.get('correction',''))[:25]}' | {str(e.get('description',''))[:50]}")

# ---------- 4. 幻觉错别字：original_text == correction 的报错 ----------
print('\n=== 幻觉错别字（original_text 与 correction 完全相同） ===')
halluc = []
for r in rows:
    errs = parse_errors(r.get('errors'))
    for e in errs:
        ot = str(e.get('original_text', ''))
        co = str(e.get('correction', ''))
        if ot and ot == co:
            halluc.append((r.get('id'), r.get('error_type'), ot[:20], r.get('TP'), r.get('FN'), r.get('FP')))
print(f'共 {len(halluc)} 处:')
for h in halluc:
    print(f'  id={h[0]} [{h[1]}] "{h[2]}" TP={h[3]} FN={h[4]} FP={h[5]}')

# ---------- 5. FP 中按描述模式的分类 ----------
print('\n=== FP 行报错的描述模式分类 ===')
fp_rows = [r for r in rows if r.get('FP')]
pattern_counter = Counter()
for r in fp_rows:
    errs = parse_errors(r.get('errors'))
    for e in errs:
        desc = str(e.get('description', ''))
        if '非选择题' in desc or '无法提取选项' in desc or '不包含选项' in desc or '无法进行选项结构' in desc:
            pattern_counter['非选择题硬套选项结构检测'] += 1
        elif '题目类型识别错误' in desc or '题型' in str(e.get('correction', '')) or '简答题' in str(e.get('correction', '')) or '选择题' in str(e.get('correction', '')) or '改写' in str(e.get('correction', '')):
            pattern_counter['题型误判/要求改写题面'] += 1
        elif '错别字或拼写错误' in desc:
            pattern_counter['疑似错别字误报'] += 1
        elif '选项标号序列' in desc or '未按A' in desc:
            pattern_counter['选项标号序列误报'] += 1
        elif '重复选项' in desc or '完全相同' in desc:
            pattern_counter['重复选项误报'] += 1
        elif '缺失' in desc or '不完整' in desc:
            pattern_counter['内容缺失类误报'] += 1
        else:
            pattern_counter['其他: ' + desc[:30]] += 1
for k, v in pattern_counter.most_common():
    print(f'  {k}: {v}')

# ---------- 6. 跳过正样本的 gold 错误内容 ----------
print('\n=== 跳过的60条正样本的 gold 错误（实际漏报） ===')
unmarked = [r for r in rows if r.get('TP') is None and r.get('TN') is None and r.get('FN') is None and r.get('FP') is None]
sk_pos = [r for r in unmarked if r.get('is_real_error') == 1]
gold_pat = Counter()
for r in sk_pos:
    gold_errs = parse_errors(r.get('correction')) if r.get('correction') else []
    for e in gold_errs:
        gold_pat[f"{r.get('error_type')}"] += 1
print(f'跳过正样本按 gold 类型: {dict(gold_pat)}')

# ---------- 7. 与训练集对比：缺失的16行 ----------
train_ids = set()
with open(r'C:\hyt-agent\训练集-微调版提示词\train_v1.1\train_v1.1_manifest.jsonl', 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        try:
            d = json.loads(line)
            # manifest 的 id 字段
            tid = d.get('id') or d.get('sample_id')
            if tid:
                train_ids.add(tid)
        except Exception:
            pass
eval_ids = set(r.get('id') for r in rows)
print(f'\n训练集 manifest id 数: {len(train_ids)}, 评估文件 id 数: {len(eval_ids)}')
if train_ids:
    missing = sorted(train_ids - eval_ids)
    print(f'缺失的 id（{len(missing)} 个）: {missing[:20]}')
