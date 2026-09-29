"""完整数据集标点符号_v2.xlsx 强制校验（需求12全部项）。"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from punct_repair import inject

from openpyxl import load_workbook

JL = r'C:\hyt-agent\训练集-微调版提示词\标点符号错误_训练集_v2.jsonl'
OUT = r'C:\hyt-agent\数据集\完整数据集标点符号_v2.xlsx'
TPL = r'C:\hyt-agent\数据集\完整数据集标点符号.xlsx'

COLUMNS = ['id', 'school', 'file_name', 'title', 'content', 'error_content',
           'error_cnt', 'error_type', 'error_detailed_type', 'correction',
           'design_reason', 'is_content_same', 'is_real_error', 'corrected_content',
           'model_name']

errors_found = []


def check(cond, msg):
    status = '✓' if cond else '✗'
    print(f'  [{status}] {msg}')
    if not cond:
        errors_found.append(msg)


# 0. 文件可以正常打开
print('== 0. 文件打开 ==')
wb = load_workbook(OUT)
ws = wb[wb.sheetnames[0]]
check(wb.sheetnames == ['Sheet1'], f"工作表名称: {wb.sheetnames}（与模板一致）")

header = [c.value for c in ws[1]]
rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
wb.close()
print(f'  打开成功，工作表 Sheet1，{len(rows)} 数据行')

# 1. 列名与顺序
print('== 1. 列名与顺序 ==')
check(header == COLUMNS, f'15 列列名与顺序完全一致: {header == COLUMNS}')

# 2. 行数
print('== 2. 行数 ==')
check(len(rows) == 1061, f'数据行数 = {len(rows)}（期望 1061）')

# 3. 有错/无错数量
print('== 3. 正负样本 ==')
n_pos = sum(1 for r in rows if r['is_real_error'])
n_neg = sum(1 for r in rows if not r['is_real_error'])
check(n_pos == 580 and n_neg == 481, f'有错 {n_pos}（期望580）| 无错 {n_neg}（期望481）')

# 4. id 连续且不重复
print('== 4. id ==')
ids = [r['id'] for r in rows]
check(all(isinstance(i, int) for i in ids), 'id 全为整数')
check(ids == list(range(1, 1062)), 'id 从1到1061连续不重复')

# 5. 每条记录都有题干
print('== 5. 题干非空 ==')
empty_ec = [r['id'] for r in rows if not r['error_content'] or not str(r['error_content']).strip()]
empty_ct = [r['id'] for r in rows if not r['content'] or not str(r['content']).strip()]
check(not empty_ec, f'error_content 无空值（空值行: {empty_ec[:5]}）')
check(not empty_ct, f'content 无空值（空值行: {empty_ct[:5]}）')

# 6. error_type 固定
print('== 6. 错误大类 ==')
bad_type = [r['id'] for r in rows if r['error_type'] != '标点符号错误']
check(not bad_type, f'error_type 全部为"标点符号错误"（异常行: {bad_type[:5]}）')

# 7. 读取 JSONL 做逐条一致性校验
print('== 7. JSONL 与 XLSX 逐条一致 ==')
jl = []
with open(JL, 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        d = json.loads(line)
        user_c = asst_c = ''
        for m in d['messages']:
            if m['role'] == 'user':
                user_c = m['content']
            elif m['role'] == 'assistant':
                asst_c = m['content']
        payload = json.loads(user_c.split('[INPUT_PAYLOAD]', 1)[1].strip())
        jl.append({'ref': payload['reference'], 'dc': payload['detection_content'], 'a': json.loads(asst_c)})
check(len(jl) == 1061, f'JSONL 记录数 = {len(jl)}')

# 7a. 题干逐条严格一致（error_content == detection_content，含换行）
mismatch = [r['id'] for r, j in zip(rows, jl) if r['error_content'] != j['dc']]
check(not mismatch, f'error_content 与 detection_content 逐条严格相等（不一致行: {mismatch[:5]}）')

# 7b. 标签一致
lab_mm = [r['id'] for r, j in zip(rows, jl) if bool(r['is_real_error']) != bool(j['a']['has_error'])]
check(not lab_mm, f'is_real_error 与 has_error 逐条一致（不一致行: {lab_mm[:5]}）')

# 7c. error_cnt 一致
cnt_mm = [r['id'] for r, j in zip(rows, jl) if r['error_cnt'] != j['a']['total_errors']]
check(not cnt_mm, f'error_cnt 与 total_errors 逐条一致（不一致行: {cnt_mm[:5]}）')

# 7d. reference 元数据一致
meta_mm = []
for r, j in zip(rows, jl):
    m_s = re.search(r'【学科】(.+)', j['ref'])
    m_t = re.search(r'【试卷标题】(.+)', j['ref'])
    m_q = re.search(r'【题型或其它信息】(.+)', j['ref'])
    sch = m_s.group(1).strip() if m_s else ''
    fn = (m_t.group(1).strip().lstrip('#').strip() + '.pdf') if m_t else ''
    ti = m_q.group(1).strip() if m_q else ''
    if r['school'] != sch or r['file_name'] != fn or r['title'] != ti:
        meta_mm.append(r['id'])
check(not meta_mm, f'school/file_name/title 与 reference 逐条一致（不一致行: {meta_mm[:5]}）')

# 8. 580 道有错题的完整标准答案
print('== 8. 有错题标准答案完整性 ==')
sub_fields = {'error_type', 'position', 'original_text', 'anchor_text', 'correction', 'description', 'suggestion'}
bad_corr = []
for r, j in zip(rows, jl):
    if not r['is_real_error']:
        if r['correction'] is not None:
            bad_corr.append((r['id'], '负样本correction非空'))
        continue
    c = r['correction']
    if not isinstance(c, str):
        bad_corr.append((r['id'], 'correction非字符串'))
        continue
    try:
        arr = json.loads(c)
    except Exception as e:
        bad_corr.append((r['id'], f'JSON解析失败:{e}'))
        continue
    if not isinstance(arr, list):
        bad_corr.append((r['id'], '非数组'))
        continue
    if len(arr) != j['a']['total_errors']:
        bad_corr.append((r['id'], '数组长度与total_errors不符'))
        continue
    for e in arr:
        missing = sub_fields - set(e.keys())
        if missing:
            bad_corr.append((r['id'], f'缺失字段{missing}'))
        if e.get('error_type') != '标点符号错误':
            bad_corr.append((r['id'], f"error_type异常:{e.get('error_type')}"))
    # JSON 内容与源 assistant errors 完全一致
    if json.loads(c) != j['a']['errors']:
        bad_corr.append((r['id'], 'JSON内容与源errors不一致'))
check(not bad_corr, f'580 有错行 correction 全部为合法完整 errors JSON（问题行: {bad_corr[:5]}）')

# 9. content 与 errors 自洽性：正向注入验证（content 上应用错误应还原 error_content）
print('== 9. content 自洽性（正向注入验证）==')
inject_fail = []
for r, j in zip(rows, jl):
    if not r['is_real_error']:
        if r['content'] != r['error_content']:
            inject_fail.append((r['id'], '负样本content与error_content不等'))
        if r['is_content_same'] is not True:
            inject_fail.append((r['id'], '负样本is_content_same非True'))
        continue
    if r['is_content_same'] is not False:
        inject_fail.append((r['id'], '正样本is_content_same非False'))
    out = r['content']
    for e in j['a']['errors']:
        out = inject(out, e)
        if out is None:
            inject_fail.append((r['id'], '正向注入定位失败'))
            break
    if out is not None and out != r['error_content']:
        inject_fail.append((r['id'], '注入结果与error_content不等'))
check(not inject_fail, f'content+errors 注入 == error_content 全部自洽（问题行: {inject_fail[:5]}）')

# 10. 数据类型检查
print('== 10. 数据类型 ==')
check(all(isinstance(r['error_cnt'], int) and r['error_cnt'] >= 0 for r in rows), 'error_cnt 全为非负整数')
check(all(isinstance(r['is_real_error'], bool) for r in rows), 'is_real_error 全为 bool')
check(all(isinstance(r['is_content_same'], bool) for r in rows), 'is_content_same 全为 bool')
check(all(isinstance(r['id'], int) for r in rows), 'id 全为整数')

# 11. 无重复行（按 error_content+id 组合）；重复题面统计
print('== 11. 重复检查 ==')
dup_ids = [i for i in ids if ids.count(i) > 1]
check(not dup_ids, f'id 无重复')
from collections import Counter
dc_counter = Counter(r['error_content'] for r in rows)
dup_dc = {k: v for k, v in dc_counter.items() if v > 1}
print(f'  (info) 题面完全相同的行组数: {len(dup_dc)}（模板本身存在同题多行，仅统计不判失败）')

# 12. 无乱码/截断（检查非法字符和异常长度）
print('== 12. 乱码/截断 ==')
def has_garbled(s):
    if s is None:
        return False
    s = str(s)
    # 常见乱码特征：替换字符、孤立代理区
    return '\ufffd' in s or any('\ud800' <= ch <= '\udfff' for ch in s)
garbled = [r['id'] for r in rows if any(has_garbled(r[k]) for k in ('content', 'error_content', 'title', 'school'))]
check(not garbled, f'无替换字符/代理区乱码（异常行: {garbled[:5]}）')
# 截断检查：error_content 与 JSONL 严格相等已保证（7a），无截断
check(not mismatch, '题干无截断（7a 已严格验证）')

# 13. 原始文件未被修改
print('== 13. 原始文件保护 ==')
jl_now = os.stat(JL)
print(f'  JSONL: {jl_now.st_size} bytes, mtime={int(jl_now.st_mtime)}')
check(jl_now.st_size == 11736891, 'JSONL 大小不变（11736891 bytes）')
tpl_now = os.stat(TPL)
print(f'  模板: {tpl_now.st_size} bytes, mtime={int(tpl_now.st_mtime)}')
check(tpl_now.st_size == 149784, '模板大小不变（149784 bytes）')

# 14. 输出文件信息
print('== 14. 输出文件 ==')
sz = os.path.getsize(OUT)
print(f'  {OUT}')
print(f'  大小: {sz} bytes')

print()
if errors_found:
    print(f'❌ 校验失败 {len(errors_found)} 项:')
    for m in errors_found:
        print('  -', m)
    raise SystemExit(1)
else:
    print('✅ 全部校验通过')
