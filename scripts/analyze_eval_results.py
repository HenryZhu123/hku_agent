# -*- coding: utf-8 -*-
"""分析 train_v1.1_智能评估.xlsx 的审校错误（FP/FN/跳过），输出详细汇总"""
import json
from collections import Counter
from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\ASUS\Downloads\train_v1.1_智能评估.xlsx', read_only=True)
ws = wb['聚合检测']

header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
idx = {name: i for i, name in enumerate(header)}

rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    rows.append(dict(zip(header, r)))

print(f'总数据行: {len(rows)}')

# ---------- 1. 总体统计 ----------
tp = sum(1 for r in rows if r.get('TP') == 1)
tn = sum(1 for r in rows if r.get('TN') == 1)
fn = sum(1 for r in rows if r.get('FN') == 1)
fp = sum(1 for r in rows if r.get('FP') == 1)
print(f'TP={tp} TN={tn} FN={fn} FP={fp}')

# 检查是否有行没有任何标记
no_mark = [r for r in rows if r.get('TP') not in (1, None) or (r.get('TP') is None and r.get('TN') is None and r.get('FN') is None and r.get('FP') is None)]
print(f'四项标记全空的行数: {sum(1 for r in rows if r.get("TP") is None and r.get("TN") is None and r.get("FN") is None and r.get("FP") is None)}')

# evaluation_status 分布
es = Counter(str(r.get('evaluation_status')) for r in rows)
print(f'evaluation_status 分布: {dict(es)}')

# task_status 分布
ts = Counter(str(r.get('task_status')) for r in rows)
print(f'task_status 分布: {dict(ts)}')

# ---------- 2. 按 error_type 分组的错误统计 ----------
print('\n=== 按错误类型分组 ===')
by_type = {}
for r in rows:
    et = r.get('error_type', '?')
    d = by_type.setdefault(et, {'total': 0, 'TP': 0, 'TN': 0, 'FN': 0, 'FP': 0, 'skip': 0, 'other': 0})
    d['total'] += 1
    if r.get('TP') == 1: d['TP'] += 1
    elif r.get('TN') == 1: d['TN'] += 1
    elif r.get('FN') == 1: d['FN'] += 1
    elif r.get('FP') == 1: d['FP'] += 1
    else: d['other'] += 1
for et, d in sorted(by_type.items()):
    print(f'{et}: 总{d["total"]} | TP={d["TP"]} TN={d["TN"]} FN={d["FN"]} FP={d["FP"]} 无标记={d["other"]}')

# ---------- 3. FN 漏报明细（全部列出） ----------
fn_rows = [r for r in rows if r.get('FN') == 1]
print(f'\n=== 漏报 FN 明细（{len(fn_rows)} 条，全部列出） ===')
for r in fn_rows:
    det_et = r.get('detect_error_type')
    he = r.get('has_error')
    reason = str(r.get('reason') or '')[:60]
    gold_cnt = r.get('error_cnt')
    # errors 模型输出
    errs = r.get('errors')
    err_brief = ''
    if errs:
        try:
            el = json.loads(errs) if isinstance(errs, str) else errs
            if el:
                err_brief = f'模型报{len(el)}条错: ' + str(el[0].get('original_text', ''))[:30] + '→' + str(el[0].get('correction', ''))[:20]
        except Exception:
            err_brief = str(errs)[:60]
    print(f"id={r.get('id')} [{r.get('error_type')}] gold错误数={gold_cnt} | has_error={he} det_type={det_et} | {err_brief}")
    print(f"    reason: {reason}")

# ---------- 4. FP 误报统计 ----------
fp_rows = [r for r in rows if r.get('FP') == 1]
print(f'\n=== 误报 FP 统计（{len(fp_rows)} 条） ===')
fp_by_type = Counter(r.get('error_type') for r in fp_rows)
print(f'按 gold error_type 分组: {dict(fp_by_type)}')

# FP 中模型报告的错误数分布
fp_cnt = Counter()
for r in fp_rows:
    errs = r.get('errors')
    n = 0
    if errs:
        try:
            el = json.loads(errs) if isinstance(errs, str) else errs
            n = len(el)
        except Exception:
            n = -1
    fp_cnt[n] += 1
print(f'FP 中模型报告的错误条数分布: {dict(sorted(fp_cnt.items(), key=lambda x: str(x[0])))}')

# ---------- 5. 无标记行（跳过98条） ----------
unmarked = [r for r in rows if r.get('TP') is None and r.get('TN') is None and r.get('FN') is None and r.get('FP') is None]
print(f'\n=== 无 TP/TN/FN/FP 标记的行（{len(unmarked)} 条） ===')
um_by_type = Counter(r.get('error_type') for r in unmarked)
print(f'按类型: {dict(um_by_type)}')
um_by_status = Counter(str(r.get('evaluation_status')) for r in unmarked)
print(f'按 evaluation_status: {dict(um_by_status)}')
um_by_fail = Counter(str(r.get('failure_type')) for r in unmarked)
print(f'按 failure_type: {dict(um_by_fail)}')
um_ids = [r.get('id') for r in unmarked]
print(f'id 列表: {um_ids}')

# ---------- 6. detect_error_type 与 error_type 不一致但仍是 TP 的（类型混淆） ----------
tp_rows = [r for r in rows if r.get('TP') == 1]
mismatch_type = [r for r in tp_rows if r.get('detect_error_type') and r.get('detect_error_type') != r.get('error_type')]
print(f'\n=== TP 中 detect_error_type 与 gold 不一致（{len(mismatch_type)} 条） ===')
for r in mismatch_type[:10]:
    print(f"id={r.get('id')} gold={r.get('error_type')} detect={r.get('detect_error_type')}")

wb.close()
