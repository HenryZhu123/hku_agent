"""核查：has_error=false 但 errors 非空的行，FP 是评估系统原判还是我自己统计的。

重点：
1. 这些行在源评估文件里评估系统自己标的 FP/TP/TN/FN 是什么
2. 这些行的 has_error 列实际值
3. has_error=False 但 errors 非空的行，被评估系统判为 FP/TP/TN/FN 的分布
"""
from openpyxl import load_workbook
from collections import Counter

MY = r'C:\hyt-agent\训练集-微调版提示词\train_v1.1\模型表现\train_v1.1_错题汇总_checkpoint-175.xlsx'
SRC = r'C:\Users\ASUS\Downloads\train_v1.1_智能评估.xlsx'

# 1. 在我生成的错题Excel（FP误报明细工作表）里找含"当前内容属于非选择题"的行
wb = load_workbook(MY, read_only=True)
print('我Excel工作表:', wb.sheetnames)
ws = wb['FP误报明细']
header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
mine = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(header, r))
    if '当前内容属于非选择题' in str(d.get('模型errors', '')):
        mine.append(d)
wb.close()
print(f'【FP误报明细】匹配"当前内容属于非选择题"的行数: {len(mine)}')
for d in mine[:3]:
    print(f'  id={d.get("id")} | 出错分类={d.get("出错分类")} | has_error={d.get("模型has_error")} | gold={d.get("题目错误类型")}')

# 2. 在源评估文件里把这些id找出来核对
wb = load_workbook(SRC, read_only=True)
ws = wb['聚合检测']
sh = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
src = {str(r[sh.index('id')]): dict(zip(sh, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[sh.index('id')]}
wb.close()

# 3. 重要：你看到的截图所在行 — 即 has_error=False 的样本被如何标记
print()
print('=== 关键问题：has_error=False 但 errors 列非空的行，评估系统怎么判？===')
contradict = []
for sid, s in src.items():
    errs = str(s.get('errors', '')).strip()
    if errs and errs != 'None' and errs != '[]':
        he = s.get('has_error')
        # has_error 为 False/0/'false' 之一，且 errors 非空
        is_false = str(he).strip().lower() in ('false', '0', '0.0', 'none') or he == 0 or he is False
        if is_false:
            contradict.append(s)

print(f'源文件中 has_error=False 且 errors 非空 的行数: {len(contradict)}')
print('这些行在评估系统里的标记分布:')
def mark(s):
    if s.get('FP'): return 'FP'
    if s.get('TP'): return 'TP'
    if s.get('TN'): return 'TN'
    if s.get('FN'): return 'FN'
    return '无标记'
mc = Counter(mark(s) for s in contradict)
for k, v in mc.most_common():
    print(f'  {k}: {v}')

# 4. 看几张截图所属行（错误类型=选项错误）的具体情况
print()
print('=== 错误类型=选项错误 且 has_error=False 的样本，评估系统判什么？===')
opt_w_false_he = [s for s in contradict if str(s.get('error_type', '')).strip() in ('选项结构错误', '选项错误')]
print(f'样本数: {len(opt_w_false_he)}')
omc = Counter(mark(s) for s in opt_w_false_he)
for k, v in omc.most_common():
    print(f'  {k}: {v}')

# 5. 看几张 FP 行的实例
print()
print('=== 几张"has_error=False 但被评估系统判为FP"的样本详情 ===')
he_false_but_fp = [s for s in contradict if s.get('FP') and not s.get('TP')]
print(f'总数: {len(he_false_but_fp)}')
for s in he_false_but_fp[:3]:
    print(f'--- id={s.get("id")} | gold={s.get("error_type")} | is_real={s.get("is_real_error")} ---')
    print(f'  has_error: {s.get("has_error")}')
    print(f'  errors: {str(s.get("errors"))[:200]}')
    print(f'  reason: {str(s.get("reason"))[:120]}')
